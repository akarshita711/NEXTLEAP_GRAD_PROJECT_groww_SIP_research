"""Score the engine (v1 and v2 prompts) against the hand-labelled golden set.
Usage: python3 evaluate.py <labelled_sheet.xlsx>   -> writes data/processed/golden_eval.md and webapp golden summary"""
import json, sys, collections as C
from openpyxl import load_workbook
from schema import batch_schema, BATCH_PROMPT_V2
from gemini import call

wb = load_workbook(sys.argv[1], data_only=True); ws = wb["Label"]
human = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]: continue
    rel = str(r[3] or "").strip().upper()
    human[r[0]] = {"text": r[2], "source": r[1], "relevant": rel == "Y", "sip_event": r[4] or "", "trigger": r[5] or "", "unmet_need": r[6] or "", "filled": rel in ("Y", "N")}
miss = [k for k, v in human.items() if not v["filled"]]
if miss: sys.exit(f"Unlabelled items: {miss}")
v1 = json.load(open("../data/processed/golden_engine_labels.json"))
ids = list(human)
items = "\n\n".join(f'[{i}] """{human[k]["text"][:1500]}"""' for i, k in enumerate(ids))
res, model = call(BATCH_PROMPT_V2.format(n=len(ids), items=items), batch_schema())
v2 = {ids[x["idx"]]: x for x in res if 0 <= x.get("idx", -1) < len(ids)}

FIELDS = ["sip_event", "trigger", "unmet_need"]
def score(engine, name):
    n = len(ids); tp = fp = fn = tn = 0; field_ok = {f: 0 for f in FIELDS}; field_n = {f: 0 for f in FIELDS}; fails = []
    for k in ids:
        h, e = human[k], engine[k]; er = bool(e["relevant"])
        if h["relevant"] and er: tp += 1
        elif h["relevant"] and not er: fn += 1; fails.append((k, "missed relevant item", h["text"][:110], "", ""))
        elif not h["relevant"] and er: fp += 1; fails.append((k, "false positive (marked relevant)", h["text"][:110], "", ""))
        else: tn += 1
        if h["relevant"]:                      # field accuracy over human-relevant items; engine marking irrelevant counts as wrong
            for f in FIELDS:
                field_n[f] += 1
                ok = er and e[f] == h[f]
                field_ok[f] += ok
                if not ok and er: fails.append((k, f"{f} mismatch", h["text"][:110], f"human={h[f]}", f"engine={e[f]}"))
    p = tp / (tp + fp) if tp + fp else 0; r = tp / (tp + fn) if tp + fn else 0
    return {"name": name, "relevance_acc": round(100 * (tp + tn) / n), "precision": round(100 * p), "recall": round(100 * r),
            "fields": {f: round(100 * field_ok[f] / field_n[f]) if field_n[f] else None for f in FIELDS}, "n_relevant_human": field_n["sip_event"], "fails": fails}
S1, S2 = score({k: v["x"] for k, v in v1.items()}, "v1 prompt (used for first 4,860 items)"), score(v2, "v2 prompt (scope fix; used by the web app)")
md = ["# Golden-set evaluation", f"Items: {len(ids)} (relevant per human: {S1['n_relevant_human']}). Model for v2 re-run: {model}.", ""]
md.append("| Prompt | Relevance accuracy | Precision | Recall | sip_event agreement | trigger agreement | unmet_need agreement |\n|---|---|---|---|---|---|---|")
for S in (S1, S2): md.append(f"| {S['name']} | {S['relevance_acc']}% | {S['precision']}% | {S['recall']}% | {S['fields']['sip_event']}% | {S['fields']['trigger']}% | {S['fields']['unmet_need']}% |")
for S in (S1, S2):
    md += ["", f"## Failures: {S['name']}", "| ID | Failure | Text (start) | Human | Engine |\n|---|---|---|---|---|"] + [f"| {a} | {b} | {c.replace('|', '/')} | {d} | {e} |" for a, b, c, d, e in S["fails"]]
open("../data/processed/golden_eval.md", "w").write("\n".join(md))
json.dump({"v1": {k: S1[k] for k in S1 if k != "fails"}, "v2": {k: S2[k] for k in S2 if k != "fails"}}, open("../data/processed/golden_eval.json", "w"), indent=1)
print("\n".join(md[:6]))
