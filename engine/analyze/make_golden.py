"""Draw a stratified random golden set (seeded) and write a BLIND Excel labelling sheet (no engine labels shown).
Engine labels are saved separately for scoring after the human has labelled."""
import json, random, sys
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill
from schema import ENUMS
random.seed(42)
rows = [json.loads(l) for l in open("../data/processed/extracted.jsonl")]
ok = [r for r in rows if r.get("x") and "[removed]" not in r["text"] and "[deleted]" not in r["text"]]
rel = [r for r in ok if r["x"]["relevant"]]; non = [r for r in ok if not r["x"]["relevant"]]
quota = {"playstore": 3, "reddit": 4, "reddit_comment": 4, "youtube": 3}          # 14 relevant, plus 6 irrelevant to test the relevance gate
pick = []
for s, n in quota.items(): pick += random.sample([r for r in rel if r["source"] == s and 30 <= len(r["text"]) <= 1200], n)
pick += random.sample([r for r in non if len(r["text"]) <= 1200], 6)
random.shuffle(pick)
key = {f"G{i+1:02d}": {"uid": r["uid"], "x": r["x"]} for i, r in enumerate(pick)}
json.dump(key, open("../data/processed/golden_engine_labels.json", "w"), indent=1, ensure_ascii=False)

wb = Workbook(); ws = wb.active; ws.title = "Label"
hdr = ["ID", "Source", "Text (read fully)", "relevant (Y/N)", "sip_event", "trigger", "unmet_need", "Notes (optional)"]
ws.append(hdr)
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
for i, r in enumerate(pick):
    ws.append([f"G{i+1:02d}", r["source"], r["text"], "", "", "", "", ""])
for col, w in zip("ABCDEFGH", (6, 14, 90, 14, 20, 26, 28, 30)): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    row[2].alignment = Alignment(wrap_text=True, vertical="top")
def dv(options, col):
    d = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True); ws.add_data_validation(d); d.add(f"{col}2:{col}{len(pick)+1}")
dv(["Y", "N"], "D"); dv(ENUMS["sip_event"], "E"); dv(ENUMS["trigger"], "F"); dv(ENUMS["unmet_need"], "G")
ws.freeze_panes = "D2"
g = wb.create_sheet("Guide")
guide = [
 ("HOW TO LABEL", "Read each text on its own. Label only what is stated or very strongly implied. When unsure choose 'unclear' / 'none'. ~5 min per 4 items. Do NOT look at any AI output."),
 ("relevant = Y", "The author describes THEIR OWN (or a specific person's) SIP / mutual-fund investing experience, expectations, doubts or behaviour. N for pure app bugs, stock/IPO/F&O trading, generic praise, promotions, unrelated Q&A."),
 ("sip_event", "What happened to their SIP: started / continuing / paused / reduced / stopped / failed_instalment / switched_platform / considering_stopping / none / unclear."),
 ("trigger", "The main stated reason behind that event or worry: cash_flow_pressure, low_or_negative_returns, expectation_mismatch (returns differ from what they expected), confusion_about_numbers, failed_payment_or_mandate, better_alternative, lost_interest_or_forgot, life_event, advice_from_others, app_issue, none, unclear."),
 ("unmet_need", "The single most important unmet need the text reveals: reassurance_whats_normal, plain_language_explanation, flexible_pause_or_reduce, failed_payment_alerts, expectation_setting, goal_progress_view, simple_returns_view, trusted_human_help, none."),
 ("If relevant = N", "Leave sip_event / trigger / unmet_need blank."),
]
for a, b in guide: g.append([a, b])
g.column_dimensions["A"].width = 22; g.column_dimensions["B"].width = 120
for row in g.iter_rows():
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
wb.save("../data/processed/golden_set_labelling_sheet.xlsx"); print("saved", len(pick))
